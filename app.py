import streamlit as st
import mysql.connector
import pandas as pd
import os
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="SI Habilitations",
    page_icon="🔐",
    layout="wide",
)

# ── CSS / palette ─────────────────────────────────────────────────────────────
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Arial');

    html, body, [class*="css"] {
        font-family: Arial, sans-serif;
    }

    /* Sidebar */
    section[data-testid="stSidebar"] {
        background-color: #042638;
    }
    section[data-testid="stSidebar"] * {
        color: #FFF5E9 !important;
    }
    section[data-testid="stSidebar"] .stSelectbox label,
    section[data-testid="stSidebar"] .stRadio label {
        color: #4CBFDC !important;
        font-weight: bold;
    }

    /* Main header */
    .main-title {
        color: #37306E;
        font-size: 2rem;
        font-weight: bold;
        border-bottom: 3px solid #EA4D49;
        padding-bottom: 0.4rem;
        margin-bottom: 1.5rem;
    }

    /* Section title */
    .section-title {
        color: #37306E;
        font-size: 1.2rem;
        font-weight: bold;
        margin-top: 1rem;
    }

    /* Cards / expanders */
    .stExpander {
        border: 1px solid #4CBFDC;
        border-radius: 6px;
    }

    /* Buttons */
    .stButton > button {
        background-color: #EA4D49;
        color: #FFF5E9;
        border: none;
        border-radius: 4px;
        font-family: Arial, sans-serif;
        font-weight: bold;
    }
    .stButton > button:hover {
        background-color: #37306E;
        color: #FFF5E9;
    }

    /* Data table header */
    thead tr th {
        background-color: #37306E !important;
        color: #FFF5E9 !important;
    }

    /* Info / success banners */
    .stAlert {
        border-radius: 4px;
    }
</style>
""", unsafe_allow_html=True)

# ── DB connection ─────────────────────────────────────────────────────────────
@st.cache_resource
def get_connection():
    return mysql.connector.connect(
        host=os.getenv("DB_HOST", "localhost"),
        user=os.getenv("DB_USER", "root"),
        password=os.getenv("DB_PASSWORD", "root"),
        database=os.getenv("DB_NAME", "si_habilitations"),
        autocommit=True,
    )


def run_query(sql, params=None, fetch=True):
    conn = get_connection()
    cur = conn.cursor(dictionary=True)
    cur.execute(sql, params or ())
    if fetch:
        return cur.fetchall()
    return None


def run_insert(sql, params=None):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute(sql, params or ())
    conn.commit()
    return cur.lastrowid


def run_update(sql, params=None):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute(sql, params or ())
    conn.commit()


# ── Helpers ───────────────────────────────────────────────────────────────────
def etablissements_list():
    return run_query("SELECT id, nom FROM etablissements ORDER BY nom")


def applications_list(id_etab=None):
    if id_etab:
        return run_query(
            "SELECT id, nom, editeur, domaine FROM applications WHERE id_etablissement=%s ORDER BY nom",
            (id_etab,),
        )
    return run_query("SELECT id, nom, editeur, domaine FROM applications ORDER BY nom")


def profils_list(id_app):
    return run_query(
        "SELECT id, nom, description FROM profils WHERE id_application=%s ORDER BY nom",
        (id_app,),
    )


def droits_list(id_app):
    return run_query(
        "SELECT id, nom, description, type_valeur FROM droits WHERE id_application=%s ORDER BY nom",
        (id_app,),
    )


def valeurs_list(id_droit):
    return run_query(
        "SELECT id, valeur, ordre FROM valeurs_liste WHERE id_droit=%s ORDER BY ordre",
        (id_droit,),
    )


def habilitation_get(id_profil, id_droit):
    rows = run_query(
        "SELECT valeur FROM habilitations WHERE id_profil=%s AND id_droit=%s",
        (id_profil, id_droit),
    )
    return rows[0]["valeur"] if rows else None


def habilitation_set(id_profil, id_droit, valeur):
    existing = run_query(
        "SELECT id FROM habilitations WHERE id_profil=%s AND id_droit=%s",
        (id_profil, id_droit),
    )
    if existing:
        run_update(
            "UPDATE habilitations SET valeur=%s WHERE id_profil=%s AND id_droit=%s",
            (valeur, id_profil, id_droit),
        )
    else:
        run_insert(
            "INSERT INTO habilitations (id_profil, id_droit, valeur) VALUES (%s,%s,%s)",
            (id_profil, id_droit, valeur),
        )


# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("## 🔐 SI Habilitations")
    st.markdown("---")

    etabs = etablissements_list()
    etab_options = {e["nom"]: e["id"] for e in etabs}
    etab_names = ["— Tous —"] + list(etab_options.keys())

    selected_etab_name = st.selectbox("Établissement", etab_names)
    selected_etab_id = etab_options.get(selected_etab_name)

    st.markdown("---")
    page = st.radio(
        "Navigation",
        ["Établissements", "Applications", "Profils", "Droits", "Matrice"],
    )

# ── Page : Établissements ─────────────────────────────────────────────────────
if page == "Établissements":
    st.markdown('<div class="main-title">Établissements</div>', unsafe_allow_html=True)

    etabs = etablissements_list()

    with st.expander("Ajouter un établissement", expanded=False):
        with st.form("add_etab"):
            nom = st.text_input("Nom")
            if st.form_submit_button("Ajouter"):
                if nom.strip():
                    run_insert("INSERT INTO etablissements (nom) VALUES (%s)", (nom.strip(),))
                    st.success("Établissement ajouté.")
                    st.rerun()
                else:
                    st.error("Le nom est obligatoire.")

    if etabs:
        st.markdown('<div class="section-title">Liste des établissements</div>', unsafe_allow_html=True)
        for e in etabs:
            with st.expander(e["nom"]):
                col1, col2 = st.columns([4, 1])
                with col1:
                    new_nom = st.text_input("Nom", value=e["nom"], key=f"etab_nom_{e['id']}")
                with col2:
                    st.write("")
                    st.write("")
                    if st.button("Modifier", key=f"etab_upd_{e['id']}"):
                        if new_nom.strip():
                            run_update(
                                "UPDATE etablissements SET nom=%s WHERE id=%s",
                                (new_nom.strip(), e["id"]),
                            )
                            st.success("Modifié.")
                            st.rerun()
                if st.button("Supprimer", key=f"etab_del_{e['id']}"):
                    run_update("DELETE FROM etablissements WHERE id=%s", (e["id"],))
                    st.warning("Établissement supprimé.")
                    st.rerun()
    else:
        st.info("Aucun établissement enregistré.")

# ── Page : Applications ───────────────────────────────────────────────────────
elif page == "Applications":
    st.markdown('<div class="main-title">Applications</div>', unsafe_allow_html=True)

    etabs = etablissements_list()
    etab_map = {e["id"]: e["nom"] for e in etabs}
    etab_opts_form = {e["nom"]: e["id"] for e in etabs}

    with st.expander("Ajouter une application", expanded=False):
        with st.form("add_app"):
            nom = st.text_input("Nom")
            editeur = st.text_input("Éditeur")
            domaine = st.text_input("Domaine")
            etab_sel = st.selectbox("Établissement", list(etab_opts_form.keys()))
            if st.form_submit_button("Ajouter"):
                if nom.strip():
                    run_insert(
                        "INSERT INTO applications (nom, editeur, domaine, id_etablissement) VALUES (%s,%s,%s,%s)",
                        (nom.strip(), editeur.strip(), domaine.strip(), etab_opts_form[etab_sel]),
                    )
                    st.success("Application ajoutée.")
                    st.rerun()
                else:
                    st.error("Le nom est obligatoire.")

    apps = applications_list(selected_etab_id)
    st.markdown('<div class="section-title">Liste des applications</div>', unsafe_allow_html=True)

    if apps:
        for app in apps:
            label = f"{app['nom']} — {etab_map.get(app.get('id_etablissement'), '?')}" if not selected_etab_id else app["nom"]
            with st.expander(label):
                col1, col2, col3, col4 = st.columns([3, 2, 2, 2])
                new_nom = col1.text_input("Nom", value=app["nom"] or "", key=f"app_nom_{app['id']}")
                new_edit = col2.text_input("Éditeur", value=app["editeur"] or "", key=f"app_edit_{app['id']}")
                new_dom = col3.text_input("Domaine", value=app["domaine"] or "", key=f"app_dom_{app['id']}")
                etab_keys = list(etab_opts_form.keys())
                cur_etab_nom = etab_map.get(app.get("id_etablissement"), etab_keys[0] if etab_keys else "")
                cur_idx = etab_keys.index(cur_etab_nom) if cur_etab_nom in etab_keys else 0
                new_etab = col4.selectbox("Établissement", etab_keys, index=cur_idx, key=f"app_etab_{app['id']}")

                c1, c2 = st.columns(2)
                if c1.button("Modifier", key=f"app_upd_{app['id']}"):
                    run_update(
                        "UPDATE applications SET nom=%s, editeur=%s, domaine=%s, id_etablissement=%s WHERE id=%s",
                        (new_nom.strip(), new_edit.strip(), new_dom.strip(), etab_opts_form[new_etab], app["id"]),
                    )
                    st.success("Modifié.")
                    st.rerun()
                if c2.button("Supprimer", key=f"app_del_{app['id']}"):
                    run_update("DELETE FROM applications WHERE id=%s", (app["id"],))
                    st.warning("Application supprimée.")
                    st.rerun()
    else:
        st.info("Aucune application pour cet établissement.")

# ── Page : Profils ────────────────────────────────────────────────────────────
elif page == "Profils":
    st.markdown('<div class="main-title">Profils</div>', unsafe_allow_html=True)

    apps = applications_list(selected_etab_id)
    if not apps:
        st.info("Sélectionnez un établissement avec des applications.")
    else:
        app_opts = {a["nom"]: a["id"] for a in apps}
        selected_app_name = st.selectbox("Application", list(app_opts.keys()))
        selected_app_id = app_opts[selected_app_name]

        with st.expander("Ajouter un profil", expanded=False):
            with st.form("add_profil"):
                nom = st.text_input("Nom du profil")
                desc = st.text_area("Description")
                if st.form_submit_button("Ajouter"):
                    if nom.strip():
                        run_insert(
                            "INSERT INTO profils (id_application, nom, description) VALUES (%s,%s,%s)",
                            (selected_app_id, nom.strip(), desc.strip()),
                        )
                        st.success("Profil ajouté.")
                        st.rerun()
                    else:
                        st.error("Le nom est obligatoire.")

        profils = profils_list(selected_app_id)
        st.markdown('<div class="section-title">Liste des profils</div>', unsafe_allow_html=True)

        if profils:
            for p in profils:
                with st.expander(p["nom"]):
                    new_nom = st.text_input("Nom", value=p["nom"], key=f"prof_nom_{p['id']}")
                    new_desc = st.text_area("Description", value=p["description"] or "", key=f"prof_desc_{p['id']}")
                    c1, c2 = st.columns(2)
                    if c1.button("Modifier", key=f"prof_upd_{p['id']}"):
                        run_update(
                            "UPDATE profils SET nom=%s, description=%s WHERE id=%s",
                            (new_nom.strip(), new_desc.strip(), p["id"]),
                        )
                        st.success("Modifié.")
                        st.rerun()
                    if c2.button("Supprimer", key=f"prof_del_{p['id']}"):
                        run_update("DELETE FROM profils WHERE id=%s", (p["id"],))
                        st.warning("Profil supprimé.")
                        st.rerun()
        else:
            st.info("Aucun profil pour cette application.")

# ── Page : Droits ─────────────────────────────────────────────────────────────
elif page == "Droits":
    st.markdown('<div class="main-title">Droits</div>', unsafe_allow_html=True)

    apps = applications_list(selected_etab_id)
    if not apps:
        st.info("Sélectionnez un établissement avec des applications.")
    else:
        app_opts = {a["nom"]: a["id"] for a in apps}
        selected_app_name = st.selectbox("Application", list(app_opts.keys()))
        selected_app_id = app_opts[selected_app_name]

        with st.expander("Ajouter un droit", expanded=False):
            with st.form("add_droit"):
                nom = st.text_input("Nom du droit")
                desc = st.text_area("Description")
                type_val = st.selectbox("Type de valeur", ["booleen", "liste"])
                valeurs_input = ""
                if type_val == "liste":
                    valeurs_input = st.text_input(
                        "Valeurs possibles (séparées par des virgules)",
                        help="Exemple : Lecture,Écriture,Admin",
                    )
                if st.form_submit_button("Ajouter"):
                    if nom.strip():
                        droit_id = run_insert(
                            "INSERT INTO droits (id_application, nom, description, type_valeur) VALUES (%s,%s,%s,%s)",
                            (selected_app_id, nom.strip(), desc.strip(), type_val),
                        )
                        if type_val == "liste" and valeurs_input.strip():
                            for i, v in enumerate(valeurs_input.split(",")):
                                v = v.strip()
                                if v:
                                    run_insert(
                                        "INSERT INTO valeurs_liste (id_droit, valeur, ordre) VALUES (%s,%s,%s)",
                                        (droit_id, v, i),
                                    )
                        st.success("Droit ajouté.")
                        st.rerun()
                    else:
                        st.error("Le nom est obligatoire.")

        droits = droits_list(selected_app_id)
        st.markdown('<div class="section-title">Liste des droits</div>', unsafe_allow_html=True)

        if droits:
            for d in droits:
                badge = "🔘 Booléen" if d["type_valeur"] == "booleen" else "📋 Liste"
                with st.expander(f"{d['nom']}  {badge}"):
                    new_nom = st.text_input("Nom", value=d["nom"], key=f"droit_nom_{d['id']}")
                    new_desc = st.text_area("Description", value=d["description"] or "", key=f"droit_desc_{d['id']}")
                    new_type = st.selectbox(
                        "Type",
                        ["booleen", "liste"],
                        index=0 if d["type_valeur"] == "booleen" else 1,
                        key=f"droit_type_{d['id']}",
                    )

                    if d["type_valeur"] == "liste":
                        st.markdown("**Valeurs de la liste**")
                        vals = valeurs_list(d["id"])
                        for v in vals:
                            vc1, vc2 = st.columns([4, 1])
                            new_val = vc1.text_input("Valeur", value=v["valeur"], key=f"val_{v['id']}")
                            if vc2.button("Suppr.", key=f"val_del_{v['id']}"):
                                run_update("DELETE FROM valeurs_liste WHERE id=%s", (v["id"],))
                                st.rerun()
                            elif new_val != v["valeur"]:
                                run_update(
                                    "UPDATE valeurs_liste SET valeur=%s WHERE id=%s",
                                    (new_val, v["id"]),
                                )

                        new_v = st.text_input("Ajouter une valeur", key=f"droit_newval_{d['id']}")
                        if st.button("Ajouter valeur", key=f"droit_addval_{d['id']}"):
                            if new_v.strip():
                                run_insert(
                                    "INSERT INTO valeurs_liste (id_droit, valeur, ordre) VALUES (%s,%s,%s)",
                                    (d["id"], new_v.strip(), len(vals)),
                                )
                                st.rerun()

                    c1, c2 = st.columns(2)
                    if c1.button("Modifier", key=f"droit_upd_{d['id']}"):
                        run_update(
                            "UPDATE droits SET nom=%s, description=%s, type_valeur=%s WHERE id=%s",
                            (new_nom.strip(), new_desc.strip(), new_type, d["id"]),
                        )
                        st.success("Modifié.")
                        st.rerun()
                    if c2.button("Supprimer", key=f"droit_del_{d['id']}"):
                        run_update("DELETE FROM droits WHERE id=%s", (d["id"],))
                        st.warning("Droit supprimé.")
                        st.rerun()
        else:
            st.info("Aucun droit pour cette application.")

# ── Page : Matrice ────────────────────────────────────────────────────────────
elif page == "Matrice":
    st.markdown('<div class="main-title">Matrice des habilitations</div>', unsafe_allow_html=True)

    apps = applications_list(selected_etab_id)
    if not apps:
        st.info("Sélectionnez un établissement avec des applications.")
    else:
        app_opts = {a["nom"]: a["id"] for a in apps}
        selected_app_name = st.selectbox("Application", list(app_opts.keys()))
        selected_app_id = app_opts[selected_app_name]

        profils = profils_list(selected_app_id)
        droits = droits_list(selected_app_id)

        if not profils:
            st.info("Aucun profil défini pour cette application.")
        elif not droits:
            st.info("Aucun droit défini pour cette application.")
        else:
            # Build valeurs cache
            valeurs_cache = {}
            for d in droits:
                if d["type_valeur"] == "liste":
                    valeurs_cache[d["id"]] = [v["valeur"] for v in valeurs_list(d["id"])]

            # Build habilitations cache
            hab_cache = {}
            for p in profils:
                for d in droits:
                    hab_cache[(p["id"], d["id"])] = habilitation_get(p["id"], d["id"])

            # Render matrix header
            header_cols = st.columns([2] + [2] * len(droits))
            header_cols[0].markdown("**Profil \\ Droit**")
            for i, d in enumerate(droits):
                type_icon = "🔘" if d["type_valeur"] == "booleen" else "📋"
                header_cols[i + 1].markdown(f"**{d['nom']}** {type_icon}")

            st.markdown("---")

            for p in profils:
                row_cols = st.columns([2] + [2] * len(droits))
                row_cols[0].markdown(f"**{p['nom']}**")

                for i, d in enumerate(droits):
                    current = hab_cache.get((p["id"], d["id"]))

                    if d["type_valeur"] == "booleen":
                        checked = current == "1" or current is True or str(current) == "True"
                        new_val = row_cols[i + 1].checkbox(
                            "",
                            value=checked,
                            key=f"hab_{p['id']}_{d['id']}",
                        )
                        db_val = "1" if new_val else "0"
                        if str(current) != db_val:
                            habilitation_set(p["id"], d["id"], db_val)

                    else:
                        vals = valeurs_cache.get(d["id"], [])
                        options = ["—"] + vals
                        cur_idx = options.index(current) if current in options else 0
                        new_val = row_cols[i + 1].selectbox(
                            "",
                            options,
                            index=cur_idx,
                            key=f"hab_{p['id']}_{d['id']}",
                        )
                        db_val = new_val if new_val != "—" else None
                        if current != db_val:
                            habilitation_set(p["id"], d["id"], db_val)

            # ── Export XLS ────────────────────────────────────────────────────
            st.markdown("---")

            def build_xlsx(app_name, profils, droits, valeurs_cache, hab_cache):
                wb = Workbook()
                ws = wb.active
                ws.title = "Matrice"

                # Styles
                fill_header = PatternFill("solid", fgColor="37306E")
                fill_profil = PatternFill("solid", fgColor="042638")
                fill_oui = PatternFill("solid", fgColor="4CBFDC")
                fill_non = PatternFill("solid", fgColor="FFF5E9")
                fill_titre = PatternFill("solid", fgColor="EA4D49")

                font_white_bold = Font(name="Arial", bold=True, color="FFFFFF")
                font_dark = Font(name="Arial", color="042638")
                font_titre = Font(name="Arial", bold=True, color="FFFFFF", size=14)

                thin = Side(style="thin", color="CCCCCC")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)
                center = Alignment(horizontal="center", vertical="center", wrap_text=True)

                # Title row
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(droits) + 1)
                title_cell = ws.cell(row=1, column=1, value=f"Matrice des habilitations — {app_name}")
                title_cell.fill = fill_titre
                title_cell.font = font_titre
                title_cell.alignment = center
                ws.row_dimensions[1].height = 30

                # Header row (droits)
                ws.cell(row=2, column=1, value="Profil \\ Droit").fill = fill_header
                ws.cell(row=2, column=1).font = font_white_bold
                ws.cell(row=2, column=1).alignment = center
                ws.cell(row=2, column=1).border = border

                for j, d in enumerate(droits):
                    cell = ws.cell(row=2, column=j + 2, value=d["nom"])
                    cell.fill = fill_header
                    cell.font = font_white_bold
                    cell.alignment = center
                    cell.border = border
                ws.row_dimensions[2].height = 40

                # Data rows
                for i, p in enumerate(profils):
                    row = i + 3
                    prof_cell = ws.cell(row=row, column=1, value=p["nom"])
                    prof_cell.fill = fill_profil
                    prof_cell.font = font_white_bold
                    prof_cell.alignment = center
                    prof_cell.border = border
                    ws.row_dimensions[row].height = 22

                    for j, d in enumerate(droits):
                        val = hab_cache.get((p["id"], d["id"]))
                        col = j + 2

                        if d["type_valeur"] == "booleen":
                            display = "Oui" if str(val) == "1" else "Non"
                            fill = fill_oui if str(val) == "1" else fill_non
                        else:
                            display = val if val else "—"
                            fill = fill_oui if val else fill_non

                        cell = ws.cell(row=row, column=col, value=display)
                        cell.fill = fill
                        cell.font = font_dark
                        cell.alignment = center
                        cell.border = border

                # Column widths
                ws.column_dimensions[get_column_letter(1)].width = 22
                for j in range(len(droits)):
                    ws.column_dimensions[get_column_letter(j + 2)].width = 16

                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                return buf.getvalue()

            xlsx_data = build_xlsx(selected_app_name, profils, droits, valeurs_cache, hab_cache)
            st.download_button(
                label="Exporter en Excel (.xlsx)",
                data=xlsx_data,
                file_name=f"matrice_{selected_app_name.replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
