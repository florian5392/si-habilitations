"""
utils/export.py — Génération des fichiers Excel (.xlsx) avec openpyxl.

Trois fonctions publiques, une pour chaque type d'export :
  - build_xlsx             : matrice complète d'une application
  - build_xlsx_comparaison : comparaison côte à côte de deux profils
  - build_xlsx_transversale: vue d'un profil sur toutes les applications

Convention openpyxl : les couleurs sont passées sans le '#' initial
(ex. "37306E" et non "#37306E").
"""

import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Palette de couleurs HospiConnect (hex sans #) ─────────────────────────────
_CORAIL    = "EA4D49"
_BLEU      = "37306E"
_BLEU_NUIT = "042638"
_TURQUOISE = "4CBFDC"
_BEIGE     = "FFF5E9"


def _fmt_val(val, type_valeur):
    """
    Formate une valeur brute de la BDD en chaîne lisible pour l'export Excel.
    Les boléens sont convertis en "Oui"/"Non" ; NULL et vide deviennent "—".
    """
    if val is None:
        return "—"
    if type_valeur == "booleen":
        return "Oui" if str(val) == "1" else "Non"
    return val if val else "—"


# ── Styles partagés ────────────────────────────────────────────────────────────

def _make_styles():
    """Construit et retourne les objets de style réutilisés dans les trois exports."""
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center")
    return {
        "fill_header":    PatternFill("solid", fgColor=_BLEU),
        "fill_profil":    PatternFill("solid", fgColor=_BLEU_NUIT),
        "fill_titre":     PatternFill("solid", fgColor=_CORAIL),
        "fill_oui":       PatternFill("solid", fgColor=_TURQUOISE),
        "fill_non":       PatternFill("solid", fgColor=_BEIGE),
        "font_white_bold": Font(name="Arial", bold=True, color="FFFFFF"),
        "font_dark":       Font(name="Arial", color=_BLEU_NUIT),
        "font_titre":      Font(name="Arial", bold=True, color="FFFFFF", size=14),
        "border": border,
        "center": center,
        "left":   left,
    }


def _save(wb):
    """Sérialise le workbook dans un BytesIO et retourne les bytes prêts à télécharger."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)   # remettre le curseur au début avant lecture
    return buf.getvalue()


# ── Export 1 : Matrice complète ────────────────────────────────────────────────

def build_xlsx(app_name, profils, droits, valeurs_cache, hab_cache):
    """
    Génère la matrice complète d'une application.

    Paramètres :
        app_name     : nom de l'application (pour le titre)
        profils      : liste de dicts {id, nom}
        droits       : liste de dicts {id, nom, type_valeur}
        valeurs_cache: dict {id_droit: [valeurs possibles]} (utilisé seulement pour info)
        hab_cache    : dict {(id_profil, id_droit): valeur} — pré-chargé pour éviter
                       N×M requêtes dans les boucles de rendu
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrice"

    s = _make_styles()

    # Titre fusionné sur toute la largeur de la matrice
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(droits) + 1)
    tc = ws.cell(row=1, column=1, value=f"Matrice des habilitations — {app_name}")
    tc.fill, tc.font, tc.alignment = s["fill_titre"], s["font_titre"], s["center"]
    ws.row_dimensions[1].height = 30

    # En-têtes : colonne 1 = label d'intersection, colonnes 2..N = noms des droits
    h = ws.cell(row=2, column=1, value="Profil \\ Droit")
    h.fill, h.font, h.alignment, h.border = s["fill_header"], s["font_white_bold"], s["center"], s["border"]
    for j, d in enumerate(droits):
        c = ws.cell(row=2, column=j + 2, value=d["nom"])
        c.fill, c.font, c.alignment, c.border = s["fill_header"], s["font_white_bold"], s["center"], s["border"]
    ws.row_dimensions[2].height = 40

    # Lignes de données : une ligne par profil, une cellule par droit
    for i, p in enumerate(profils):
        row = i + 3
        pc = ws.cell(row=row, column=1, value=p["nom"])
        pc.fill, pc.font, pc.alignment, pc.border = s["fill_profil"], s["font_white_bold"], s["center"], s["border"]
        ws.row_dimensions[row].height = 22

        for j, d in enumerate(droits):
            val = hab_cache.get((p["id"], d["id"]))
            if d["type_valeur"] == "booleen":
                display = "Oui" if str(val) == "1" else "Non"
                fill    = s["fill_oui"] if str(val) == "1" else s["fill_non"]
            else:
                display = val if val else "—"
                fill    = s["fill_oui"] if val else s["fill_non"]
            cell = ws.cell(row=row, column=j + 2, value=display)
            cell.fill, cell.font, cell.alignment, cell.border = fill, s["font_dark"], s["center"], s["border"]

    # Largeurs de colonnes : 22 pour les profils, 16 pour chaque droit
    ws.column_dimensions[get_column_letter(1)].width = 22
    for j in range(len(droits)):
        ws.column_dimensions[get_column_letter(j + 2)].width = 16

    return _save(wb)


# ── Export 2 : Comparaison de deux profils ─────────────────────────────────────

def build_xlsx_comparaison(app_name, profil_1, profil_2, rows):
    """
    Génère le fichier de comparaison côte à côte de deux profils.

    Couleurs des cellules selon le statut de chaque droit :
        Identique → bleu clair (B8EDF5)
        Différent → orange clair (FFD9A0)
        Manquant  → rouge clair (F9A8A7)

    Paramètres :
        rows : liste de dicts retournés par get_habilitations_comparaison()
               {id_droit, droit_nom, type_valeur, valeur_1, valeur_2}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparaison"

    s = _make_styles()
    fill_identique = PatternFill("solid", fgColor="B8EDF5")
    fill_different = PatternFill("solid", fgColor="FFD9A0")
    fill_manquant  = PatternFill("solid", fgColor="F9A8A7")

    # Titre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    tc = ws.cell(row=1, column=1, value=f"Comparaison — {app_name} — {profil_1} vs {profil_2}")
    tc.fill, tc.font, tc.alignment = s["fill_titre"], s["font_titre"], s["center"]
    ws.row_dimensions[1].height = 30

    # En-têtes
    for j, h in enumerate(["Droit", profil_1, profil_2, "Statut"]):
        c = ws.cell(row=2, column=j + 1, value=h)
        c.fill, c.font, c.alignment, c.border = s["fill_header"], s["font_white_bold"], s["center"], s["border"]
    ws.row_dimensions[2].height = 28

    # Données
    for i, row in enumerate(rows):
        r  = i + 3
        v1 = _fmt_val(row["valeur_1"], row["type_valeur"])
        v2 = _fmt_val(row["valeur_2"], row["type_valeur"])
        has_1 = row["valeur_1"] is not None
        has_2 = row["valeur_2"] is not None

        # Logique de classement : présence asymétrique > valeurs différentes > identiques
        if has_1 != has_2:
            fill, statut = fill_manquant, "Manquant"
        elif v1 == v2:
            fill, statut = fill_identique, "Identique"
        else:
            fill, statut = fill_different, "Différent"

        c = ws.cell(row=r, column=1, value=row["droit_nom"])
        c.fill, c.font, c.alignment, c.border = s["fill_profil"], s["font_white_bold"], s["center"], s["border"]

        for j, val in enumerate([v1, v2, statut]):
            c = ws.cell(row=r, column=j + 2, value=val)
            c.fill, c.font, c.alignment, c.border = fill, s["font_dark"], s["center"], s["border"]

        ws.row_dimensions[r].height = 20

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14

    return _save(wb)


# ── Export 3 : Vue transversale ────────────────────────────────────────────────

def build_xlsx_transversale(profil_nom, etab_nom, rows):
    """
    Génère la vue transversale d'un profil sur toutes les applications.

    Les lignes sont groupées par application avec des couleurs de fond alternées
    (bleu clair / beige) pour faciliter la lecture visuelle.

    Paramètres :
        rows : liste de dicts retournés par get_habilitations_transversales()
               {id_app, app_nom, droit_nom, type_valeur, valeur}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Vue transversale"

    s = _make_styles()
    fill_row_a  = PatternFill("solid", fgColor="EBF5FB")  # bleu très clair, applications paires
    fill_row_b  = PatternFill("solid", fgColor=_BEIGE)    # beige, applications impaires
    fill_val_ok = PatternFill("solid", fgColor=_TURQUOISE) # valeur renseignée
    fill_val_no = PatternFill("solid", fgColor="F0F0F0")   # valeur absente

    # Titre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    tc = ws.cell(row=1, column=1, value=f"Vue transversale — {profil_nom} — {etab_nom}")
    tc.fill, tc.font, tc.alignment = s["fill_titre"], s["font_titre"], s["center"]
    ws.row_dimensions[1].height = 30

    # En-têtes
    for j, h in enumerate(["Application", "Droit", "Valeur"]):
        c = ws.cell(row=2, column=j + 1, value=h)
        c.fill, c.font, c.alignment, c.border = s["fill_header"], s["font_white_bold"], s["center"], s["border"]
    ws.row_dimensions[2].height = 28

    # Données avec alternance de couleur par groupe d'application
    current_app = None
    app_idx     = 0
    fill_row    = fill_row_a

    for i, row in enumerate(rows):
        r = i + 3
        # Changement de groupe → on bascule la couleur de fond
        if row["app_nom"] != current_app:
            current_app = row["app_nom"]
            fill_row = fill_row_a if app_idx % 2 == 0 else fill_row_b
            app_idx += 1

        val    = _fmt_val(row["valeur"], row["type_valeur"])
        # Valeur considérée "renseignée" si non NULL, non vide et non "0" (booléen faux)
        is_set = row["valeur"] is not None and row["valeur"] not in ("", "0")
        fv     = fill_val_ok if is_set else fill_val_no

        c = ws.cell(row=r, column=1, value=row["app_nom"])
        c.fill, c.font, c.alignment, c.border = fill_row, s["font_dark"], s["left"], s["border"]

        c = ws.cell(row=r, column=2, value=row["droit_nom"])
        c.fill, c.font, c.alignment, c.border = fill_row, s["font_dark"], s["left"], s["border"]

        c = ws.cell(row=r, column=3, value=val)
        c.fill, c.font, c.alignment, c.border = fv, s["font_dark"], s["center"], s["border"]

        ws.row_dimensions[r].height = 20

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 18

    return _save(wb)
